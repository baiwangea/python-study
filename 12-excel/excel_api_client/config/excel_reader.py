"""
Excel配置读取器 - 从Excel文件中读取所有配置信息

支持读取以下工作表：
- API接口配置
- 请求参数配置
- 认证信息配置
- 环境配置
- 日志配置
"""

from pathlib import Path
from typing import Optional, List, Dict, Any

from openpyxl import load_workbook

from excel_api_client.models.api_config import (
    APIEndpoint, APIParameter, AuthConfig, EnvironmentConfig,
    LogConfig, ClientConfig, HTTPMethod, ParamLocation, AuthType
)
from excel_api_client.utils.logger import get_logger


logger = get_logger()


class ExcelConfigReader:
    """Excel配置读取器"""

    # 工作表名称映射
    SHEET_API_ENDPOINTS = 'API接口配置'
    SHEET_PARAMETERS = '请求参数配置'
    SHEET_AUTH = '认证信息配置'
    SHEET_ENVIRONMENT = '环境配置'
    SHEET_LOG = '日志配置'

    def __init__(self, excel_path: str):
        """
        初始化读取器
        
        Args:
            excel_path: Excel文件路径
        """
        self.excel_path = Path(excel_path)
        if not self.excel_path.exists():
            raise FileNotFoundError(f"配置文件不存在: {excel_path}")
        
        self.wb = None
        self._load_workbook()

    def _load_workbook(self) -> None:
        """加载Excel工作簿"""
        try:
            self.wb = load_workbook(self.excel_path, data_only=True)
            logger.logger.info(f"成功加载配置文件: {self.excel_path}")
            logger.logger.info(f"包含工作表: {self.wb.sheetnames}")
        except Exception as e:
            logger.log_error(e, "加载Excel文件失败")
            raise

    def _read_sheet_data(self, sheet_name: str, header_row: int = 1) -> List[Dict[str, Any]]:
        """
        读取工作表数据为字典列表
        
        Args:
            sheet_name: 工作表名称
            header_row: 表头所在行号
            
        Returns:
            每行数据为一个字典，键为表头，值为单元格内容
        """
        if sheet_name not in self.wb.sheetnames:
            logger.logger.warning(f"工作表 '{sheet_name}' 不存在，跳过")
            return []

        ws = self.wb[sheet_name]
        headers = []
        data = []

        # 读取表头
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=header_row, column=col).value
            if cell_value:
                headers.append(str(cell_value).strip())
            else:
                headers.append(f"Column_{col}")

        # 读取数据行
        for row in range(header_row + 1, ws.max_row + 1):
            row_data = {}
            has_data = False
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col_idx)
                value = cell.value
                
                # 处理空值
                if value is None or str(value).strip() == '':
                    row_data[header] = None
                else:
                    row_data[header] = value
                    has_data = True

            if has_data:
                data.append(row_data)

        logger.logger.debug(f"工作表 '{sheet_name}' 读取了 {len(data)} 行数据")
        return data

    def read_api_endpoints(self) -> List[APIEndpoint]:
        """读取API端点配置"""
        data = self._read_sheet_data(self.SHEET_API_ENDPOINTS)
        endpoints = []
        
        for row in data:
            endpoint_name = row.get('接口名称', '')
            try:
                endpoint = APIEndpoint(
                    name=endpoint_name,
                    method=row.get('请求方法', 'GET'),
                    base_url=row.get('基础URL', ''),
                    path=row.get('接口路径', ''),
                    timeout=row.get('超时时间(秒)', 10),
                    retry_count=row.get('重试次数', 3),
                    enabled=row.get('启用状态', '是'),
                    description=row.get('备注', '')
                )
                endpoints.append(endpoint)
            except Exception as e:
                logger.logger.warning(f"解析端点配置失败: 接口名称={endpoint_name or '未知'} | 错误: {e}")
                continue

        logger.logger.info(f"读取了 {len(endpoints)} 个API端点配置")
        return endpoints

    def read_parameters(self) -> List[APIParameter]:
        """读取请求参数配置"""
        data = self._read_sheet_data(self.SHEET_PARAMETERS)
        parameters = []
        
        for row in data:
            param_name = row.get('参数名称', '')
            endpoint_name = row.get('所属接口', '')
            try:
                param = APIParameter(
                    name=param_name,
                    endpoint_name=endpoint_name,
                    location=row.get('参数位置', 'query'),
                    param_type=row.get('参数类型', 'string'),
                    default_value=row.get('默认值'),
                    required=row.get('是否必填', '否'),
                    description=row.get('参数说明', '')
                )
                parameters.append(param)
            except Exception as e:
                logger.logger.warning(
                    f"解析参数配置失败: 参数名称={param_name or '未知'}, 所属接口={endpoint_name or '未知'} | 错误: {e}"
                )
                continue

        logger.logger.info(f"读取了 {len(parameters)} 个参数配置")
        return parameters

    def read_auth_configs(self) -> List[AuthConfig]:
        """读取认证信息配置"""
        data = self._read_sheet_data(self.SHEET_AUTH)
        auth_configs = []
        
        for row in data:
            environment = row.get('环境', 'dev')
            auth_type = row.get('认证类型', 'Bearer')
            try:
                auth = AuthConfig(
                    environment=environment,
                    auth_type=auth_type,
                    token=row.get('Token/Key', ''),
                    username=row.get('用户名', ''),
                    password=row.get('密码', ''),
                    expire_time=row.get('过期时间'),
                    refresh_url=row.get('刷新URL', ''),
                    description=row.get('备注', '')
                )
                auth_configs.append(auth)
            except Exception as e:
                logger.logger.warning(
                    f"解析认证配置失败: 环境={environment or '未知'}, 认证类型={auth_type or '未知'} | 错误: {e}"
                )
                continue

        logger.logger.info(f"读取了 {len(auth_configs)} 个认证配置")
        return auth_configs

    def read_environment_configs(self) -> List[EnvironmentConfig]:
        """读取环境配置"""
        data = self._read_sheet_data(self.SHEET_ENVIRONMENT)
        env_configs = []
        
        for row in data:
            config_key = row.get('配置项', '')
            try:
                env = EnvironmentConfig(
                    key=config_key,
                    current_env=row.get('当前环境', 'dev'),
                    dev_value=row.get('开发环境值', ''),
                    test_value=row.get('测试环境值', ''),
                    staging_value=row.get('预发布环境值', ''),
                    prod_value=row.get('生产环境值', ''),
                    description=row.get('配置说明', '')
                )
                env_configs.append(env)
            except Exception as e:
                logger.logger.warning(f"解析环境配置失败: 配置项={config_key or '未知'} | 错误: {e}")
                continue

        logger.logger.info(f"读取了 {len(env_configs)} 个环境配置")
        return env_configs

    def read_log_config(self) -> LogConfig:
        """读取日志配置"""
        data = self._read_sheet_data(self.SHEET_LOG)
        
        # 转换为键值对
        config_dict = {}
        for row in data:
            key = row.get('配置项', '')
            value = row.get('配置值', '')
            if key:
                config_dict[key] = value

        # 映射到LogConfig字段
        field_mapping = {
            '日志目录': 'log_dir',
            '日志文件名': 'log_file',
            '错误日志文件名': 'error_file',
            '日志格式': 'log_format',
            '日期格式': 'date_format',
            '控制台输出': 'console_output',
            '文件最大大小(MB)': 'max_file_size_mb',
            '保留日志文件数': 'backup_count',
            '请求日志记录': 'log_request',
            '响应日志记录': 'log_response',
            'Excel日志记录': 'excel_log_enabled',
            'Excel日志文件': 'excel_log_file',
            'Excel日志工作表': 'excel_log_sheet',
        }

        kwargs = {}
        for cn_key, en_key in field_mapping.items():
            if cn_key in config_dict:
                kwargs[en_key] = config_dict[cn_key]

        try:
            log_config = LogConfig(**kwargs)
            logger.logger.info("日志配置读取成功")
            return log_config
        except Exception as e:
            logger.log_error(e, "解析日志配置失败，使用默认配置")
            return LogConfig()

    def read_all(self) -> ClientConfig:
        """
        读取所有配置
        
        Returns:
            完整的客户端配置对象
        """
        endpoints = self.read_api_endpoints()
        parameters = self.read_parameters()
        auth_configs = self.read_auth_configs()
        env_configs = self.read_environment_configs()
        log_config = self.read_log_config()

        # 从环境配置中获取当前环境
        current_env = 'dev'
        for cfg in env_configs:
            if cfg.key == 'CURRENT_ENV':
                current_env = cfg.current_env or 'dev'
                break

        client_config = ClientConfig(
            endpoints=endpoints,
            parameters=parameters,
            auth_configs=auth_configs,
            env_configs=env_configs,
            log_config=log_config,
            current_environment=current_env
        )

        # 记录配置摘要
        config_summary = {
            '当前环境': current_env,
            'API端点数': len(endpoints),
            '参数配置数': len(parameters),
            '认证配置数': len(auth_configs),
            '环境配置数': len(env_configs),
            '启用端点': sum(1 for ep in endpoints if ep.enabled),
        }
        logger.log_config_loaded(config_summary)

        return client_config

    def close(self) -> None:
        """关闭工作簿"""
        if self.wb:
            self.wb.close()
            logger.logger.info("Excel工作簿已关闭")

    def __enter__(self):
        """上下文管理器入口"""
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        """上下文管理器出口"""
        self.close()
        return False
