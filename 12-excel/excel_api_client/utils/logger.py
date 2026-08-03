"""
日志工具模块 - 提供结构化的日志记录功能

支持：
- 控制台和文件双输出
- 日志轮转（按文件大小）
- 请求/响应独立记录
- 错误日志分离
"""

import copy
import logging
import logging.handlers
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook

from excel_api_client.models.api_config import LogConfig


class ColoredFormatter(logging.Formatter):
    """带颜色的日志格式化器"""
    
    COLORS = {
        'DEBUG': '\033[36m',      # 青色
        'INFO': '\033[32m',       # 绿色
        'WARNING': '\033[33m',    # 黄色
        'ERROR': '\033[31m',      # 红色
        'CRITICAL': '\033[35m',   # 紫色
    }
    RESET = '\033[0m'

    def format(self, record: logging.LogRecord) -> str:
        """格式化日志记录"""
        log_color = self.COLORS.get(record.levelname, self.RESET)
        record.levelname = f"{log_color}{record.levelname}{self.RESET}"
        return super().format(record)


class APILogger:
    """API客户端专用日志管理器"""

    _instance: Optional['APILogger'] = None
    _initialized: bool = False
    SENSITIVE_HEADERS = {'authorization', 'x-api-key', 'proxy-authorization'}
    SENSITIVE_FIELDS = {
        'token', 'access_token', 'refresh_token', 'id_token',
        'password', 'passwd', 'pwd',
        'secret', 'client_secret',
        'api_key', 'apikey', 'x_api_key'
    }

    def __new__(cls, *args, **kwargs):
        if cls._instance is None:
            cls._instance = super().__new__(cls)
        return cls._instance

    def __init__(self, config: Optional[LogConfig] = None):
        new_config = config or LogConfig()
        if APILogger._initialized:
            if config and self.config.to_dict() != new_config.to_dict():
                self.config = new_config
                self._setup_log_dir()
                self._create_loggers()
            return

        self.config = new_config
        self._setup_log_dir()
        self._create_loggers()
        APILogger._initialized = True

    def _setup_log_dir(self) -> None:
        """创建日志目录"""
        log_dir = Path(self.config.log_dir)
        log_dir.mkdir(parents=True, exist_ok=True)
        self.log_dir = log_dir

        excel_log_file = Path(self.config.excel_log_file)
        if not excel_log_file.is_absolute():
            excel_log_file = Path.cwd() / excel_log_file
        self.excel_log_file = excel_log_file

    def _create_loggers(self) -> None:
        """创建并配置日志记录器"""
        # 主日志记录器
        self.logger = logging.getLogger('api_client')
        self.logger.setLevel(logging.DEBUG)
        self.logger.handlers = []  # 清除已有handler
        self.logger.propagate = False

        # 格式化器
        file_formatter = logging.Formatter(
            self.config.log_format,
            datefmt=self.config.date_format
        )
        console_formatter = ColoredFormatter(
            self.config.log_format,
            datefmt=self.config.date_format
        )

        # 主日志文件Handler (轮转)
        main_log_path = self.log_dir / self.config.log_file
        file_handler = logging.handlers.RotatingFileHandler(
            filename=main_log_path,
            maxBytes=self.config.max_file_size_mb * 1024 * 1024,
            backupCount=self.config.backup_count,
            encoding='utf-8'
        )
        file_handler.setLevel(logging.DEBUG)
        file_handler.setFormatter(file_formatter)
        self.logger.addHandler(file_handler)

        # 错误日志文件Handler
        error_log_path = self.log_dir / self.config.error_file
        error_handler = logging.handlers.RotatingFileHandler(
            filename=error_log_path,
            maxBytes=self.config.max_file_size_mb * 1024 * 1024,
            backupCount=self.config.backup_count,
            encoding='utf-8'
        )
        error_handler.setLevel(logging.ERROR)
        error_handler.setFormatter(file_formatter)
        self.logger.addHandler(error_handler)

        # 控制台Handler
        if self.config.console_output:
            console_handler = logging.StreamHandler(sys.stdout)
            console_handler.setLevel(logging.DEBUG)
            console_handler.setFormatter(console_formatter)
            self.logger.addHandler(console_handler)

        # 请求/响应专用记录器
        self.request_logger = logging.getLogger('api_client.request')
        self.request_logger.setLevel(logging.DEBUG)
        self.request_logger.handlers = []
        self.request_logger.propagate = False

        request_log_path = self.log_dir / 'api_requests.log'
        request_handler = logging.handlers.RotatingFileHandler(
            filename=request_log_path,
            maxBytes=self.config.max_file_size_mb * 1024 * 1024,
            backupCount=self.config.backup_count,
            encoding='utf-8'
        )
        request_handler.setFormatter(file_formatter)
        self.request_logger.addHandler(request_handler)

        if self.config.console_output:
            req_console = logging.StreamHandler(sys.stdout)
            req_console.setFormatter(console_formatter)
            self.request_logger.addHandler(req_console)

    def _mask_value(self, value: Any) -> str:
        """统一的脱敏占位值"""
        if value is None:
            return '***'
        value_str = str(value)
        if not value_str:
            return '***'
        return '***REDACTED***'

    def _sanitize_headers(self, headers: Optional[dict]) -> dict:
        """脱敏请求/响应头"""
        sanitized = copy.deepcopy(headers or {})
        for key in list(sanitized.keys()):
            if str(key).lower() in self.SENSITIVE_HEADERS:
                sanitized[key] = self._mask_value(sanitized[key])
        return sanitized

    def _sanitize_data(self, data: Any) -> Any:
        """递归脱敏常见敏感字段"""
        if isinstance(data, dict):
            sanitized = copy.deepcopy(data)
            for key, value in sanitized.items():
                if str(key).lower() in self.SENSITIVE_FIELDS:
                    sanitized[key] = self._mask_value(value)
                else:
                    sanitized[key] = self._sanitize_data(value)
            return sanitized
        if isinstance(data, list):
            return [self._sanitize_data(item) for item in data]
        if isinstance(data, tuple):
            return tuple(self._sanitize_data(item) for item in data)
        return data

    def _sanitize_text(self, text: str) -> str:
        """脱敏文本中的常见敏感键值"""
        sanitized_text = text
        patterns = [
            r'(?i)(authorization\s*[:=]\s*)([^,\s\]\}]+(?:\s+[^,\s\]\}]+)?)',
            r'(?i)(x-api-key\s*[:=]\s*)([^,\s\]\}]+)',
            r'(?i)("(?:token|access_token|refresh_token|id_token|password|passwd|pwd|secret|client_secret|api_key|apikey|x_api_key)"\s*:\s*")([^"]*)(")',
            r"(?i)('(?:token|access_token|refresh_token|id_token|password|passwd|pwd|secret|client_secret|api_key|apikey|x_api_key)'\s*:\s*')([^']*)(')",
            r'(?i)((?:token|access_token|refresh_token|id_token|password|passwd|pwd|secret|client_secret|api_key|apikey|x_api_key)\s*[:=]\s*)([^,\s\]\}]+)'
        ]
        for pattern in patterns[:2]:
            sanitized_text = re.sub(pattern, lambda m: f"{m.group(1)}{self._mask_value(m.group(2))}", sanitized_text)
        for pattern in patterns[2:4]:
            sanitized_text = re.sub(pattern, lambda m: f"{m.group(1)}{self._mask_value(m.group(2))}{m.group(3)}", sanitized_text)
        sanitized_text = re.sub(
            patterns[4],
            lambda m: f"{m.group(1)}{self._mask_value(m.group(2))}",
            sanitized_text
        )
        return sanitized_text

    def _prepare_log_payload(self, data: Any) -> str:
        """将日志数据转换为脱敏后的字符串"""
        sanitized_data = self._sanitize_data(data)
        return self._sanitize_text(str(sanitized_data))

    def _append_excel_log(self, log_type: str, method_or_status: str, url_or_duration: str, content: str) -> None:
        """将日志追加到Excel工作表"""
        if not self.config.excel_log_enabled:
            return
        try:
            if not self.excel_log_file.exists():
                return
            wb = load_workbook(self.excel_log_file)
            sheet_name = self.config.excel_log_sheet
            if sheet_name not in wb.sheetnames:
                ws = wb.create_sheet(sheet_name)
                ws.append(['时间', '类型', '方法/状态', 'URL/耗时', '内容'])
            else:
                ws = wb[sheet_name]

            ws.append([
                datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                log_type,
                method_or_status,
                url_or_duration,
                content,
            ])
            wb.save(self.excel_log_file)
            wb.close()
        except Exception as exc:
            self.logger.warning(f"写入Excel日志失败: {exc}")

    def log_request(self, method: str, url: str, headers: dict, params: Optional[dict] = None,
                    body: Optional[dict] = None) -> None:
        """记录请求信息"""
        if not self.config.log_request:
            return
        
        safe_headers = self._sanitize_headers(headers)
        safe_params = self._prepare_log_payload(params) if params else None
        safe_body = self._prepare_log_payload(body) if body else None

        self.request_logger.info("=" * 60)
        self.request_logger.info(f"[REQUEST] {method} {url}")
        self.request_logger.info(f"[HEADERS] {safe_headers}")
        content_parts = [f"HEADERS: {safe_headers}"]
        if safe_params:
            self.request_logger.info(f"[PARAMS]  {safe_params}")
            content_parts.append(f"PARAMS: {safe_params}")
        if safe_body:
            self.request_logger.info(f"[BODY]    {safe_body}")
            content_parts.append(f"BODY: {safe_body}")
        self.request_logger.info("-" * 60)
        self._append_excel_log('REQUEST', method, url, ' | '.join(content_parts))

    def log_response(self, status_code: int, response_headers: dict, 
                     response_body: Optional[str] = None, duration_ms: Optional[float] = None) -> None:
        """记录响应信息"""
        if not self.config.log_response:
            return
        
        duration_str = f" ({duration_ms:.2f}ms)" if duration_ms else ""
        safe_headers = self._sanitize_headers(response_headers)
        self.request_logger.info(f"[RESPONSE] Status: {status_code}{duration_str}")
        self.request_logger.info(f"[RESPONSE HEADERS] {safe_headers}")
        content_parts = [f"HEADERS: {safe_headers}"]
        if response_body:
            body_str = self._prepare_log_payload(response_body)
            if len(body_str) > 2000:
                body_str = body_str[:2000] + f"... [截断，总长度: {len(body_str)}]"
            self.request_logger.info(f"[RESPONSE BODY] {body_str}")
            content_parts.append(f"BODY: {body_str}")
        self.request_logger.info("=" * 60)
        self._append_excel_log('RESPONSE', str(status_code), duration_str.strip() or '-', ' | '.join(content_parts))

    def log_error(self, error: Exception, context: Optional[str] = None) -> None:
        """记录错误信息"""
        msg = f"{context}: {str(error)}" if context else str(error)
        self.logger.error(msg, exc_info=True)

    def log_retry(self, attempt: int, max_retries: int, error: Exception) -> None:
        """记录重试信息"""
        self.logger.warning(
            f"请求失败，正在进行第 {attempt}/{max_retries} 次重试 | 错误: {str(error)}"
        )

    def log_config_loaded(self, config_summary: dict) -> None:
        """记录配置加载信息"""
        self.logger.info("=" * 60)
        self.logger.info("配置加载完成")
        for key, value in config_summary.items():
            self.logger.info(f"  {key}: {value}")
        self.logger.info("=" * 60)


def get_logger(config: Optional[LogConfig] = None) -> APILogger:
    """获取日志管理器实例"""
    return APILogger(config)
